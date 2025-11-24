from __future__ import annotations
import io
import os
import re
import fitz  # PyMuPDF
import pdfplumber
import pytesseract
from PIL import Image
from bs4 import BeautifulSoup
from docx import Document
from pptx import Presentation
from typing import Tuple
from app.utils.log_utils import log_error


def extract_pdf(path: str) -> str:
    try:
        with pdfplumber.open(path) as pdf:
            text = ''.join(page.extract_text() or '' for page in pdf.pages)
        if len(text) < 100:
            text += _ocr_pdf(path)
        return text
    except Exception as exc:  # noqa: BLE001
        log_error(f"PDF extraction failed for {path}: {exc}")
        return ''


def _ocr_pdf(path: str) -> str:
    content = ''
    try:
        doc = fitz.open(path)
        for page in doc:
            pix = page.get_pixmap()
            img = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
            content += pytesseract.image_to_string(img)
    except Exception as exc:  # noqa: BLE001
        log_error(f"PDF OCR failed for {path}: {exc}")
    return content


def extract_docx(path: str) -> str:
    try:
        doc = Document(path)
        return '\n'.join(p.text for p in doc.paragraphs)
    except Exception as exc:  # noqa: BLE001
        log_error(f"DOCX extraction failed for {path}: {exc}")
        return ''


def extract_pptx(path: str) -> str:
    try:
        prs = Presentation(path)
        parts = []
        for slide in prs.slides:
            if slide.shapes.title:
                parts.append(slide.shapes.title.text)
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    parts.append(shape.text)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                parts.append(slide.notes_slide.notes_text_frame.text)
        return '\n'.join(parts)
    except Exception as exc:  # noqa: BLE001
        log_error(f"PPTX extraction failed for {path}: {exc}")
        return ''


def extract_image(path: str) -> str:
    try:
        img = Image.open(path)
        return pytesseract.image_to_string(img)
    except Exception as exc:  # noqa: BLE001
        log_error(f"Image OCR failed for {path}: {exc}")
        return ''


def extract_html(path: str) -> Tuple[str, list[str]]:
    try:
        with open(path, 'r', encoding='utf-8', errors='ignore') as f:
            html = f.read()
        soup = BeautifulSoup(html, 'html.parser')
        for tag in soup(['script', 'style']):
            tag.decompose()
        links = [a.get('href') for a in soup.find_all('a', href=True)]
        text = soup.get_text(separator='\n')
        return text, links
    except Exception as exc:  # noqa: BLE001
        log_error(f"HTML extraction failed for {path}: {exc}")
        return '', []


def extract_text_generic(path: str) -> Tuple[str, list[str]]:
    ext = os.path.splitext(path)[1].lower()
    links: list[str] = []
    text = ''
    if ext == '.pdf':
        text = extract_pdf(path)
    elif ext == '.docx':
        text = extract_docx(path)
    elif ext == '.pptx':
        text = extract_pptx(path)
    elif ext in {'.png', '.jpg', '.jpeg', '.bmp', '.tiff'}:
        text = extract_image(path)
    elif ext in {'.html', '.htm'}:
        text, links = extract_html(path)
    else:
        try:
            with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
        except Exception as exc:  # noqa: BLE001
            log_error(f"Fallback text extraction failed for {path}: {exc}")
            text = ''
    links += re.findall(r'https?://\S+', text)
    return text, links
